#!/usr/bin/env python3
"""Export confirmed level IDs from DashForge's 1,820-slot workbook.

Only rows with a non-empty level ID are exported. The report keeps track of
empty slots; this script never fabricates IDs or fills missing historical cells.
"""
from __future__ import annotations
import argparse
from pathlib import Path
from openpyxl import load_workbook

def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", type=Path, required=True)
    ap.add_argument("--sheet", default="Matriz 1820")
    ap.add_argument("--output", type=Path, default=Path("corpus-level-ids.txt"))
    ap.add_argument("--id-column", type=int, default=7, help="1-based ID column; default 7")
    args = ap.parse_args()
    if not args.workbook.exists():
        ap.error(f"Workbook not found: {args.workbook}")
    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    if args.sheet not in wb.sheetnames:
        ap.error(f"Sheet {args.sheet!r} not found. Available: {', '.join(wb.sheetnames)}")
    ws = wb[args.sheet]
    ids = []
    empty = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        value = row[args.id_column - 1] if len(row) >= args.id_column else None
        if value is None or str(value).strip() == "":
            empty += 1
            continue
        level_id = str(value).strip()
        if level_id.isdigit():
            ids.append(level_id)
    unique = sorted(set(ids), key=int)
    args.output.write_text("\n".join(unique) + ("\n" if unique else ""), encoding="utf-8")
    print(f"Unique numeric level IDs exported: {len(unique)}")
    print(f"Rows without an ID: {empty}")
    print(f"Saved: {args.output}")
    if empty:
        print("WARNING: workbook still has empty IDs; these cannot be used for training.")
    if len(unique) < 1820:
        print(f"WARNING: only {len(unique)} unique IDs for a target of 1,820; corpus is incomplete.")
    return 0

if __name__ == "__main__":
    raise SystemExit(main())
